"""DOCX -> блоки в порядке документа"""

import re
from collections.abc import Iterator
from pathlib import Path

from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from edu_doc_agent.parsing.models import BaseParser, Block, BlockKind, CellSpan, ParsedDocument

_TEXT_XPATH = (
    ".//w:t[not(ancestor::*[local-name()='Fallback'])] | .//m:t | .//w:r/w:tab | .//w:r/w:br"
)
_HEADING_STYLE = re.compile(r"^(?:heading|заголовок)\s*(\d)$", re.IGNORECASE)
_SECTION = re.compile(r"^раздел\s+\d", re.IGNORECASE)  # раздел
_NUMBERED = re.compile(r"^(\d+(?:\.\d+)+)\.?\s")  # нумерация
_BULLET = re.compile(r"^[•·▪◦‣−–—-]\s+")  # маркер
_BLANK = re.compile(r"_{4,}")  # поле для заполнения '______' -> '...'
_RULE = re.compile(r"[\s_.…\-–—=*]{4,}")  # абзац целиком из таких символов - линия-разделитель
_SPACES = re.compile(r"[ \t\u00a0]+")  # пробелы, табы, неразрывные пробелы
_MAX_HEADING_LEN = 200


class DocxParser(BaseParser):
    def parse(self, path: Path) -> ParsedDocument:
        doc = Document(str(path))
        numbering = _Numbering(doc)
        blocks = []
        for item in _iter_body(doc, doc.element.body):
            if isinstance(item, Table):
                block = _table(item)
            else:
                block = _paragraph(item, numbering)
            if block is not None:
                blocks.append(block)
        return ParsedDocument(source=str(path), blocks=blocks)


def _iter_body(doc: DocxDocument, element) -> Iterator[Paragraph | Table]:
    """Абзацы и таблицы основного текста по порядку"""
    for child in element.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _iter_body(doc, content)


# Абзацы:


def _paragraph(p: Paragraph, numbering: "_Numbering") -> Block | None:
    number, list_level, is_section = numbering.next(p)
    raw = _text(p._p)
    if _RULE.fullmatch(raw.strip()):  # '--------', '________', '........' без текста
        return None
    text = _clean(raw)
    if not text:
        return None
    style = p.style.name if p.style is not None else ""
    is_bold = _is_bold(p)
    is_title = not text.endswith((".", ";", ","))
    is_short = len(text) <= _MAX_HEADING_LEN and not text.endswith((";", ","))

    if match := _HEADING_STYLE.match(style):
        return Block(BlockKind.HEADING, text, level=int(match.group(1)), number=number)
    if list_level is not None:
        if number and ((is_section and is_title) or (is_bold and is_short)):
            return Block(BlockKind.HEADING, text, level=list_level + 1, number=number)
        return Block(BlockKind.LIST_ITEM, text, level=list_level, number=number)
    typed_section = _SECTION.match(text) or _NUMBERED.match(text)
    if is_bold and (is_short or (typed_section and is_title)):
        return Block(BlockKind.HEADING, text, level=_heading_level(text))
    if match := _BULLET.match(text):
        return Block(BlockKind.LIST_ITEM, text[match.end() :])
    if style.startswith("List"):
        return Block(BlockKind.LIST_ITEM, text)
    return Block(BlockKind.PARAGRAPH, text)


def _text(element) -> str:
    parts = []
    for node in element.xpath(_TEXT_XPATH):
        name = node.tag.rsplit("}", 1)[-1]
        if name == "t":
            parts.append(node.text or "")
        else:
            parts.append("\n" if name == "br" else " ")
    return "".join(parts)


def _clean(text: str, keep_newlines: bool = False) -> str:
    text = _BLANK.sub("…", text.replace("\xad", ""))
    lines = [_SPACES.sub(" ", line).strip() for line in text.splitlines()]
    return ("\n" if keep_newlines else " ").join(line for line in lines if line)


def _is_bold(p: Paragraph) -> bool:
    runs = [r for r in p.runs if r.text.strip()]
    return bool(runs) and all(r.bold for r in runs)


def _heading_level(text: str) -> int:
    if _SECTION.match(text):
        return 1
    if match := _NUMBERED.match(text):
        return match.group(1).count(".") + 1
    return 3


class _Numbering:
    def __init__(self, doc: DocxDocument) -> None:
        self.levels: dict[tuple[str, int], tuple[str, str]] = {}
        self.counters: dict[str, list[int]] = {}
        try:
            root = doc.part.numbering_part.element
        except (KeyError, NotImplementedError):
            return
        abstract = {a.get(qn("w:abstractNumId")): a for a in root.findall(qn("w:abstractNum"))}
        for num in root.findall(qn("w:num")):
            ref = num.find(qn("w:abstractNumId"))
            definition = abstract.get(ref.get(qn("w:val"))) if ref is not None else None
            if definition is None:
                continue
            for lvl in definition.findall(qn("w:lvl")):
                fmt, template = lvl.find(qn("w:numFmt")), lvl.find(qn("w:lvlText"))
                self.levels[(num.get(qn("w:numId")), int(lvl.get(qn("w:ilvl"))))] = (
                    fmt.get(qn("w:val")) if fmt is not None else "decimal",
                    template.get(qn("w:val")) if template is not None else "",
                )

    def next(self, p: Paragraph) -> tuple[str, int | None, bool]:
        num_pr = _num_pr(p)
        if num_pr is None or num_pr.numId is None or num_pr.numId.val == 0:
            return "", None, False
        num_id = str(num_pr.numId.val)
        level = num_pr.ilvl.val if num_pr.ilvl is not None else 0
        counters = self.counters.setdefault(num_id, [0] * 9)
        counters[level] += 1
        counters[level + 1 :] = [0] * (8 - level)

        fmt, template = self.levels.get((num_id, level), ("bullet", ""))
        if fmt in ("bullet", "none"):
            return "", level, False

        def number(match: re.Match[str]) -> str:
            k = int(match.group(1)) - 1
            k_fmt = self.levels.get((num_id, k), ("decimal", ""))[0]
            return _format_number(max(counters[k], 1), k_fmt)

        is_section = fmt == "decimal" and template.count("%") > 1
        return re.sub(r"%(\d)", number, template).strip(), level, is_section


def _num_pr(p: Paragraph):
    for ppr in (p._p.pPr, p.style.element.pPr if p.style is not None else None):
        if ppr is not None and ppr.numPr is not None:
            return ppr.numPr
    return None


_RU_LETTERS = "абвгдежзиклмнопрстуфхцчшщэюя"
_ROMAN = {1000: "m", 900: "cm", 500: "d", 400: "cd", 100: "c", 90: "xc", 50: "l", 40: "xl"}
_ROMAN |= {10: "x", 9: "ix", 5: "v", 4: "iv", 1: "i"}


def _format_number(n: int, fmt: str) -> str:
    if fmt in ("lowerLetter", "upperLetter"):
        s = chr(ord("a") + (n - 1) % 26)
    elif fmt in ("russianLower", "russianUpper"):
        s = _RU_LETTERS[(n - 1) % len(_RU_LETTERS)]
    elif fmt in ("lowerRoman", "upperRoman"):
        s = ""
        for value, digits in _ROMAN.items():
            count, n = divmod(n, value)
            s += digits * count
    else:
        return str(n)
    return s.upper() if fmt.startswith("upper") or fmt == "russianUpper" else s


# Таблицы:


def _table(table: Table) -> Block | None:
    grid = []  # [(элемент ячейки, текст)] по строкам; пустые строки пропускаем
    for row in table.rows:
        cells = row.cells
        line = [
            (cell._tc, "" if i > 0 and cell._tc is cells[i - 1]._tc else _cell_text(cell))
            for i, cell in enumerate(cells)
        ]
        if any(text for _, text in line):
            grid.append(line)
    if not grid:
        return None
    width = max(len(line) for line in grid)
    rows = [[text for _, text in line] + [""] * (width - len(line)) for line in grid]
    return Block(BlockKind.TABLE, rows=rows, spans=_spans(grid))


def _spans(grid: list) -> list[CellSpan]:
    """Объединенные ячейки: позиции сетки, которые занимает один и тот же элемент ячейки"""
    positions: dict[int, list[tuple[int, int]]] = {}
    for r, line in enumerate(grid):
        for c, (tc, _) in enumerate(line):
            positions.setdefault(id(tc), []).append((r, c))
    spans = []
    for cells in positions.values():
        rows, cols = {r for r, _ in cells}, {c for _, c in cells}
        if len(cells) > 1:
            spans.append(CellSpan(min(rows), min(cols), len(rows), len(cols)))
    return sorted(spans, key=lambda s: (s.row, s.col))


def _cell_text(cell) -> str:
    return _clean("\n".join(_text(p._p) for p in cell.paragraphs), keep_newlines=True)
