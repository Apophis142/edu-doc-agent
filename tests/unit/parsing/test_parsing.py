"""Тесты Parser Service на синтетическом DOCX"""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from edu_doc_agent.parsing import (
    Block,
    BlockKind,
    CellSpan,
    ParsedDocument,
    UnsupportedFormatError,
    parse_document,
)

# формула Word 'x+1=0' и 'элемент управления содержимым'
_FORMULA = (
    f"<w:p {nsdecls('w', 'm')}><w:r><w:t xml:space='preserve'>Решите </w:t></w:r>"
    "<m:oMath><m:r><m:t>x+1=0</m:t></m:r></m:oMath></w:p>"
)
_CONTENT_CONTROL = (
    f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r><w:t>Задание</w:t></w:r></w:p>"
    "</w:sdtContent></w:sdt>"
)
# многоуровневая автонумерация '1.»'/ '1.1.'
_NUMBERING = (
    f"<w:abstractNum {nsdecls('w')} w:abstractNumId='90'>"
    "<w:lvl w:ilvl='0'><w:numFmt w:val='decimal'/><w:lvlText w:val='%1.'/></w:lvl>"
    "<w:lvl w:ilvl='1'><w:numFmt w:val='decimal'/><w:lvlText w:val='%1.%2.'/></w:lvl>"
    "</w:abstractNum>"
)
_NUM = f"<w:num {nsdecls('w')} w:numId='90'><w:abstractNumId w:val='90'/></w:num>"


@pytest.fixture
def rpd(tmp_path: Path) -> Path:
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "КОЛОНТИТУЛ"
    doc.add_heading("Раздел 1. Характеристики", level=1)  # стиль Heading 1
    doc.add_paragraph().add_run("1.2. Вид практики").bold = True  # жирный с номером
    doc.add_paragraph().add_run("Аннотация").bold = True  # жирный без номера
    doc.add_paragraph("Целью\xa0курса   является ...")
    doc.add_paragraph("изучить материал лекции;", style="List Bullet")
    doc.add_paragraph("   ")  # пустой
    body = doc.element.body
    body.insert(len(body) - 1, parse_xml(_FORMULA))
    body.insert(len(body) - 1, parse_xml(_CONTENT_CONTROL))

    t = doc.add_table(rows=5, cols=3)
    t.cell(0, 0).text = "Компетенция"
    t.cell(0, 1).text = "Результаты"
    t.cell(0, 1).merge(t.cell(0, 2))  # шапка на две колонки
    for i, v in enumerate(["ОПК-2", "ОПК-2.1", "Знать: формы"]):
        t.cell(1, i).text = v
    t.cell(2, 1).text = "ОПК-2.2"
    t.cell(2, 2).text = "Уметь: SQL"
    t.cell(1, 0).merge(t.cell(2, 0))  # ОПК-2 на две строки
    t.cell(3, 0).text = "ПК-1"
    t.cell(3, 1).text = "ПК-1.3"
    t.cell(3, 2).text = "Первая строка"
    t.cell(3, 2).add_paragraph("вторая | строка")
    path = tmp_path / "rpd.docx"
    doc.save(path)
    return path


@pytest.fixture
def characteristic(tmp_path: Path) -> Path:
    doc = Document()
    numbering = doc.part.numbering_part.element
    numbering.insert(0, parse_xml(_NUMBERING))
    numbering.append(parse_xml(_NUM))

    def numbered(text: str, level: int, bold: bool = False) -> None:
        p = doc.add_paragraph()
        p.add_run(text).bold = bold
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = level
        num_pr.get_or_add_numId().val = 90

    numbered("Характеристика деятельности", 0, bold=True)
    numbered("Квалификация:", 1)
    doc.add_paragraph("Бакалавр")
    numbered("Области деятельности", 1)
    doc.add_paragraph("знать модель;", style="List Number")
    doc.add_paragraph("уметь писать SQL.", style="List Number")
    doc.add_paragraph("от ____________ № __________")
    doc.add_paragraph("________________________________")
    doc.add_paragraph("– пункт с маркером, набранным вручную")
    doc.add_paragraph("* Примечание к таблице")
    path = tmp_path / "characteristic.docx"
    doc.save(path)
    return path


def test_blocks(rpd: Path) -> None:
    blocks = parse_document(rpd).blocks
    assert [(b.kind, b.level, b.text) for b in blocks[:-1]] == [
        (BlockKind.HEADING, 1, "Раздел 1. Характеристики"),
        (BlockKind.HEADING, 2, "1.2. Вид практики"),
        (BlockKind.HEADING, 3, "Аннотация"),
        (BlockKind.PARAGRAPH, 0, "Целью курса является ..."),
        (BlockKind.LIST_ITEM, 0, "изучить материал лекции;"),
        (BlockKind.PARAGRAPH, 0, "Решите x+1=0"),  # формула не потерялась
        (BlockKind.PARAGRAPH, 0, "Задание"),
    ]
    assert blocks[-1].rows == [
        ["Компетенция", "Результаты", ""],  # объединение по горизонтали -> пусто
        ["ОПК-2", "ОПК-2.1", "Знать: формы"],
        ["ОПК-2", "ОПК-2.2", "Уметь: SQL"],  # объединение по вертикали -> повтор
        ["ПК-1", "ПК-1.3", "Первая строка\nвторая | строка"],
    ]


def test_merged_cells_are_described_by_spans(tmp_path: Path) -> None:
    """Как таблица 2.1 РПД: двухуровневая шапка с объединениями и служебная строка."""
    doc = Document()
    t = doc.add_table(rows=4, cols=4)
    t.cell(0, 0).text = "Период обучения"
    t.cell(0, 0).merge(t.cell(1, 0))  # по вертикали — две строки шапки
    t.cell(0, 1).text = "Контактная работа"
    t.cell(0, 1).merge(t.cell(0, 2))  # по горизонтали — группа над двумя столбцами
    t.cell(1, 1).text = "лекции"
    t.cell(1, 2).text = "семинары"
    t.cell(0, 3).text = "Трудоёмкость"
    t.cell(0, 3).merge(t.cell(1, 3))
    t.cell(2, 0).text = "ОСНОВНАЯ ТРАЕКТОРИЯ"
    t.cell(2, 0).merge(t.cell(2, 3))  # служебная строка на всю ширину
    for i, value in enumerate(["Семестр 1", "30", "30", "3"]):
        t.cell(3, i).text = value
    path = tmp_path / "hours.docx"
    doc.save(path)

    (table,) = parse_document(path).tables
    assert table.rows == [  # плоская сетка (и Markdown) — без изменений
        ["Период обучения", "Контактная работа", "", "Трудоёмкость"],
        ["Период обучения", "лекции", "семинары", "Трудоёмкость"],
        ["ОСНОВНАЯ ТРАЕКТОРИЯ", "", "", ""],
        ["Семестр 1", "30", "30", "3"],
    ]
    assert table.spans == [
        CellSpan(row=0, col=0, row_span=2),
        CellSpan(row=0, col=1, col_span=2),
        CellSpan(row=0, col=3, row_span=2),
        CellSpan(row=2, col=0, col_span=4),
    ]
    spans = parse_document(path).to_dict()["blocks"][0]["spans"]
    assert spans[0] == {"row": 0, "col": 0, "row_span": 2, "col_span": 1}  # так видит extraction


MARKDOWN_CONTRACT = """\
# Раздел 1. Характеристики

## 1.2. Вид практики

### Аннотация

Целью курса является ...

- изучить материал лекции;

Решите x+1=0

Задание

| Компетенция | Результаты |  |
| --- | --- | --- |
| ОПК-2 | ОПК-2.1 | Знать: формы |
| ОПК-2 | ОПК-2.2 | Уметь: SQL |
| ПК-1 | ПК-1.3 | Первая строка<br>вторая \\| строка |
"""


def test_markdown_contract(rpd: Path) -> None:
    """Формат Markdown — контракт с модулем извлечения (см. docstring parsing/models.py).

    Если тест упал из-за намеренного изменения формата — обновите контракт и предупредите
    тех, кто работает с extraction: от этого формата зависят их регулярные выражения.
    """
    assert parse_document(rpd).to_markdown() == MARKDOWN_CONTRACT


def test_numbering_and_cleanup(characteristic: Path) -> None:
    expected = (
        "# 1. Характеристика деятельности\n\n"
        "## 1.1. Квалификация:\n\n"
        "Бакалавр\n\n"
        "## 1.2. Области деятельности\n\n"
        "1. знать модель;\n\n"
        "2. уметь писать SQL.\n\n"
        "от … № …\n\n"
        "- пункт с маркером, набранным вручную\n\n"
        "\\* Примечание к таблице\n"
    )
    assert parse_document(characteristic).to_markdown() == expected


def test_filler_cleanup_keeps_regular_text(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Подпись ____________ (ФИО)")
    doc.add_paragraph("См. пункты 1.1–1.4 .... и далее ---- без изменений")
    doc.add_paragraph("--------")
    doc.add_paragraph("........")
    t = doc.add_table(rows=1, cols=3)
    for i, value in enumerate(["----", "....", "______"]):
        t.cell(0, i).text = value
    path = tmp_path / "form.docx"
    doc.save(path)

    blocks = parse_document(path).blocks
    assert [b.text for b in blocks[:-1]] == [
        "Подпись … (ФИО)",  # поле для заполнения -> '…'
        "См. пункты 1.1–1.4 .... и далее ---- без изменений",  # обычный текст не трогаем
    ]  # абзацы-линии '--------' и '........' убраны
    assert blocks[-1].rows == [["----", "....", "…"]]  # в таблицах '----' и '....' остаются


def test_list_indent_is_never_deeper_than_parent() -> None:
    doc = ParsedDocument(
        source="x",
        blocks=[
            Block(BlockKind.PARAGRAPH, "Перечень:"),
            Block(BlockKind.LIST_ITEM, "без родителя", level=1),
            Block(BlockKind.LIST_ITEM, "вложенный", level=1),
            Block(BlockKind.LIST_ITEM, "вложенный глубже", level=3),
        ],
    )
    assert doc.to_markdown() == (
        "Перечень:\n\n- без родителя\n\n   - вложенный\n\n      - вложенный глубже\n"
    )


def test_unsupported_format(tmp_path: Path) -> None:
    with pytest.raises(UnsupportedFormatError):
        parse_document(tmp_path / "scan.pdf")


def test_documents_to_markdown_files(rpd: Path, characteristic: Path, tmp_path: Path) -> None:
    out = tmp_path / "out"
    out.mkdir()
    for path in sorted(tmp_path.glob("*.docx")):
        markdown = parse_document(path).to_markdown()
        (out / f"{path.stem}.md").write_text(markdown, encoding="utf-8")
    assert sorted(p.name for p in out.iterdir()) == ["characteristic.md", "rpd.md"]
    assert (out / "rpd.md").read_text(encoding="utf-8").startswith("# Раздел 1.")
