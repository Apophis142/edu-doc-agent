"""Все сущности из ADR-0002 («Какие сущности извлекаем») и замечание заказчика про таблицу 2.1.

Тесты идут по строкам таблицы из ADR. Сущности из учебного плана (xlsx) и эталонные списки
компетенций из характеристики программы и матрицы лежат в других документах и здесь не
проверяются.
"""

from pathlib import Path

import pytest
from docx import Document

from edu_doc_agent.extraction import RPDExtraction, extract_rpd, extract_rpd_file
from edu_doc_agent.parsing import parse_document

FIXTURES = Path(__file__).resolve().parents[1] / "fixtures"


@pytest.fixture(scope="module")
def rpd() -> RPDExtraction:
    return extract_rpd_file(FIXTURES / "rpd_full_synthetic.md")


@pytest.fixture(scope="module")
def rpp() -> RPDExtraction:
    return extract_rpd_file(FIXTURES / "rpp_synthetic.md")


# --- MVP: РПД -----------------------------------------------------------------------------


def test_competences_indicators_outcomes_and_kim(rpd: RPDExtraction) -> None:
    """РПД 1.3 и 3.1.4: компетенции, индикаторы, результаты обучения, КИМ по индикаторам."""
    assert [c.code for c in rpd.competences] == ["УК-1"]
    assert [(i.code, i.level) for i in rpd.indicators] == [
        ("УК-1.1", "Базовый"),
        ("УК-1.2", "Продвинутый"),
    ]
    assert {(o.kind, o.text.rstrip(".")) for o in rpd.learning_outcomes} == {
        ("знает", "основные виды тестов"),
        ("умеет", "строить тестовые наборы"),
        ("владеет", "навыками анализа покрытия"),
    }
    assert rpd.assessments[0].indicator_code == "УК-1.1"
    assert rpd.assessments[0].materials == ["Контрольная работа № 1", "вопросы к экзамену"]


def test_control_forms_with_several_kinds(rpd: RPDExtraction) -> None:
    """Формы контроля по семестрам, в том числе «зачёт, экзамен»."""
    (form,) = rpd.control_forms
    assert form.semester == 1
    assert form.current == ["контрольные работы", "коллоквиум"]
    assert [a.kind for a in form.intermediate] == ["зачёт", "экзамен"]


def test_workload_by_kind_and_semester(rpd: RPDExtraction) -> None:
    """Таблица 2.1: часы по видам занятий, по семестрам, з.е., ИТОГО; без строки «1-60»."""
    sem, total = rpd.workload_table
    assert (sem.semester, sem.is_total) == (1, False)
    assert sem.contact_hours == {
        "лекции": 32,
        "консультации": 2,
        "практические занятия": 32,
        "контрольные работы": 4,
        "коллоквиумы": 2,
        "промежуточная аттестация": 2,
    }
    assert sem.self_study_hours == {"с использованием методических материалов": 34}
    assert (sem.credits, sem.total_hours, sem.active_hours) == (3, 108, 34)
    assert sem.study_form == "очная"
    assert total.is_total and total.credits == 3
    assert not [w for w in rpd.warnings if w.startswith("трудоёмкость")]


def test_identification(rpd: RPDExtraction) -> None:
    """Титул: название, рег. номер, тип, код дисциплины, язык обучения."""
    assert rpd.title == "Теория тестовых данных"
    assert rpd.title_en == "Test Data Theory"
    assert rpd.doc_kind == "discipline"
    assert rpd.registration_number == "009999"
    assert rpd.discipline_code == "009999"
    assert rpd.language == "русский"
    assert rpd.credits == 3


def test_active_forms_hours(rpd: RPDExtraction) -> None:
    """РПД 1.4: «… – 32 часа», «… – 2 часа»; предложение без часов не берём."""
    assert [f.hours for f in rpd.active_forms] == [32, 2]
    assert rpd.active_forms[0].text.startswith("Все практические занятия")
    assert rpd.active_forms[0].source.line is not None


def test_active_hours_are_comparable_with_table_column(rpd: RPDExtraction) -> None:
    """Сумма часов из 1.4 равна столбцу «Объём активных и интерактивных форм» из 2.1."""
    assert sum(f.hours or 0 for f in rpd.active_forms) == rpd.workload_table[0].active_hours


def test_prerequisites_as_graph_edges(rpd: RPDExtraction) -> None:
    """РПД 1.2: названия дисциплин в кавычках и область знаний."""
    disciplines = [p.name for p in rpd.prerequisites if p.kind == "дисциплина"]
    assert disciplines == ["Алгебра и геометрия", "Основы дискретной математики"]
    assert any(p.kind == "область знаний" for p in rpd.prerequisites)


def test_assessment_means_counts(rpd: RPDExtraction) -> None:
    """РПД 3.1.3: число контрольных и коллоквиумов."""
    counts = {m.kind: m.count for m in rpd.control_mentions if m.role == "текущий"}
    assert counts["контрольная работа"] == 2
    assert counts["коллоквиум"] == 1


def test_assessment_means_ticket_structure(rpd: RPDExtraction) -> None:
    """РПД 3.1.3: структура билета «два теоретических вопроса и одна задача»."""
    (ticket,) = rpd.tickets
    assert [(p.count, p.item) for p in ticket.parts] == [
        (2, "теоретических вопроса"),
        (1, "задачу"),
    ]


def test_assessment_means_grade_scales(rpd: RPDExtraction) -> None:
    """РПД 3.1.3: шкала оценивания A–F для зачёта и баллы билета для экзамена."""
    credit = [r for r in rpd.grade_scales if r.attestation == "зачёт"]
    exam = [r for r in rpd.grade_scales if r.attestation == "экзамен"]
    assert [(r.grade, r.ects) for r in credit] == [
        ("зачтено", "A"),
        ("зачтено", "B"),
        ("не зачтено", "F"),
    ]
    assert [(r.grade, r.ects, r.criteria) for r in exam] == [
        ("отлично", "A", "9-10"),
        ("неудовлетворительно", "F", "менее 3"),
    ]


# --- «Далее» -----------------------------------------------------------------------------


def test_topic_hours_match_table_2_1(rpd: RPDExtraction) -> None:
    """РПД 2.2: часы по темам, семестр берётся из строки над таблицей; сверка с 2.1."""
    assert len(rpd.topic_hours) == 5
    first = rpd.topic_hours[0]
    assert (first.semester, first.number, first.topic, first.kind, first.hours) == (
        1,
        "1",
        "Виды тестов",
        "лекции",
        16,
    )
    assert not [w for w in rpd.warnings if "2.2" in w]


def test_topic_hours_mismatch_with_2_1_gives_warning() -> None:
    text = (FIXTURES / "rpd_full_synthetic.md").read_text(encoding="utf-8")
    broken = extract_rpd(
        text.replace("| 2 | Покрытие | лекции | 16 |", "| 2 | Покрытие | лекции | 20 |"), "x.md"
    )
    assert any(
        "часы по темам (2.2), семестр 1, лекции: 36 ч, в таблице 2.1 — 32" in w
        for w in broken.warnings
    )


def test_exam_questions_and_example_tasks(rpd: RPDExtraction) -> None:
    """РПД 3.1.4: перечень вопросов к экзамену и примеры задач."""
    by_kind = {m.kind: m for m in rpd.control_materials}
    assert by_kind["вопросы к экзамену"].items == [
        "Виды тестов.",
        "Критерии покрытия.",
        "Мутационное тестирование.",
    ]
    assert by_kind["примеры задач"].items[0] == "Построить тестовый набор для функции abs."
    assert len(by_kind["примеры задач"].items) == 2


def test_direction_profile(rpd: RPDExtraction) -> None:
    """Профиль берётся из строки траектории таблицы; направления в этой РПД нет."""
    assert rpd.program is not None
    assert rpd.program.profiles == ["ТОП-ИТ ПРО"]
    assert rpd.program.direction_code is None


def test_direction_and_program_from_text() -> None:
    text = "Направление: 02.03.02 «Фундаментальная информатика»\n\nООП: Программирование\n"
    program = extract_rpd(text, "x.md").program
    assert program is not None
    assert (program.direction_code, program.direction, program.program) == (
        "02.03.02",
        "Фундаментальная информатика",
        "Программирование",
    )


def test_goals_and_developers(rpd: RPDExtraction) -> None:
    """РПД 1.1 и раздел 4."""
    assert rpd.goals is not None and "Целью освоения дисциплины" in rpd.goals
    assert "Задача курса" in rpd.goals
    (dev,) = rpd.developers
    assert (dev.name, dev.degree, dev.title, dev.position, dev.contacts) == (
        "Иванов Иван Иванович",
        "к.ф.-м.н.",
        "доцент",
        "доцент",
        "i.ivanov@spbu.ru",
    )


# --- РПП ---------------------------------------------------------------------------------


def test_practice_kind_type_method(rpp: RPDExtraction) -> None:
    """РПП 1.2, 1.2.1, 1.3: вид (отмеченный пункт), тип и способ без вложенных пунктов 1.3.1."""
    assert rpp.doc_kind == "practice"
    assert rpp.practice is not None
    assert (rpp.practice.kind, rpp.practice.type, rpp.practice.method) == (
        "Учебная",
        "проектная",
        "стационарная",
    )


def test_practice_workload_by_semester_from_text(rpp: RPDExtraction) -> None:
    """РПП 2.1: таблицы нет, з.е. и часы по семестрам перечислены текстом."""
    assert rpp.credits == 12
    assert [(w.semester, w.credits, w.hours) for w in rpp.workload if w.semester] == [
        (2, 6, 216),
        (3, 6, 216),
    ]
    assert rpp.semesters == [2, 3]


def test_prof_standards_and_functions(rpp: RPDExtraction) -> None:
    """РПП 1.6: профстандарты и обобщённые трудовые функции."""
    assert [(s.code, s.name) for s in rpp.prof_standards] == [
        ("06.001", "Программист"),
        ("06.003", "Архитектор программного обеспечения"),
    ]
    assert [f.split()[0] for f in rpp.prof_standards[0].functions] == ["A.3", "B.4"]


def test_competences_by_group(rpp: RPDExtraction) -> None:
    """РПП 1.7: компетенции из списков, группа — УК, УНК, УКБ, ОПК, ПК."""
    assert {c.code: c.group for c in rpp.competences} == {
        "УК-1": "УК",
        "ПКП-УНК-01": "УНК",
        "УКБ-1": "УКБ",
        "ОПК-1": "ОПК",
        "ПКА-2": "ПК",
    }


def test_competence_to_standard_links(rpp: RPDExtraction) -> None:
    """РПП 1.8: компетенции → трудовые функции по профстандартам."""
    (link,) = rpp.standard_links
    assert link.competence_codes == ["ПКА-2", "ПКП-УНК-01"]
    assert link.functions == {"06.003": ["A.6", "B.7"], "06.001": ["A.3"]}


# --- замечание заказчика: таблица 2.1 с объединёнными ячейками ----------------------------

_LEAVES = [
    "Период обучения (модуль)",
    "лекции",
    "семинары",
    "консультации",
    "практические занятия",
    "лабораторные работы",
    "контрольные работы",
    "коллоквиумы",
    "текущий контроль",
    "промежуточная аттестация",
    "итоговая аттестация",
    "под руководством преподавателя",
    "в присутствии преподавателя",
    "сам.раб. с использованием методических материалов",
    "текущий контроль (сам.раб.)",
    "промежуточная аттестация (сам.раб.)",
    "итоговая аттестация (сам.раб.)",
    "Объём активных и интерактивных форм учебных занятий",
    "Трудоёмкость",
]
_SEM = {1: 62, 3: 2, 4: 60, 6: 4, 7: 2, 9: 4, 13: 51, 15: 31, 17: 62, 18: 6}  # колонка → часы


def _docx_with_merged_table(path: Path) -> None:
    """Таблица 2.1 так, как она устроена в реальных РПД: шапка в 2 строки, объединения в обе
    стороны, служебные строки на всю ширину (траектория, форма обучения)."""
    doc = Document()
    doc.add_heading("2.1. Организация учебных занятий", level=2)
    table = doc.add_table(rows=8, cols=19)
    top = table.cell(0, 0).merge(table.cell(0, 18))
    top.text = "Трудоёмкость, объёмы учебной работы и наполняемость групп обучающихся"
    table.cell(1, 0).merge(table.cell(2, 0)).text = _LEAVES[0]
    table.cell(1, 1).merge(
        table.cell(1, 12)
    ).text = "Контактная работа обучающихся с преподавателем"
    table.cell(1, 13).merge(table.cell(1, 16)).text = "Самостоятельная работа"
    table.cell(1, 17).merge(table.cell(2, 17)).text = _LEAVES[17]
    table.cell(1, 18).merge(table.cell(2, 18)).text = _LEAVES[18]
    for col in range(1, 17):
        table.cell(2, col).text = _LEAVES[col]
    table.cell(3, 0).merge(table.cell(3, 18)).text = "ОСНОВНАЯ ТРАЕКТОРИЯ"
    table.cell(4, 0).merge(table.cell(4, 18)).text = "очная форма обучения"
    table.cell(5, 0).text = "Семестр 1"
    table.cell(6, 0).text = "Семестр 1"
    table.cell(7, 0).text = "ИТОГО"
    for col, hours in _SEM.items():
        table.cell(5, col).text = str(hours)
        table.cell(7, col).text = str(hours)
    for col in (1, 3, 4, 6, 7, 9):
        table.cell(6, col).text = "1-20"  # размер группы, не часы
    doc.save(path)


def test_merged_cells_table_2_1_from_docx(tmp_path: Path) -> None:
    """docx → parser (объединённые ячейки) → извлечение: часы по видам занятий в нужных столбцах."""
    path = tmp_path / "rpd.docx"
    _docx_with_merged_table(path)
    parsed = parse_document(path)
    assert parsed.tables[0].spans  # парсер передал объединения

    result = extract_rpd(parsed.to_markdown(), file="rpd.docx")
    sem, total = result.workload_table
    assert (sem.semester, sem.study_form, sem.trajectory) == (1, "очная", "ОСНОВНАЯ ТРАЕКТОРИЯ")
    assert sem.contact_hours == {
        "лекции": 62,
        "консультации": 2,
        "практические занятия": 60,
        "контрольные работы": 4,
        "коллоквиумы": 2,
        "промежуточная аттестация": 4,
    }
    assert sem.self_study_hours == {
        "с использованием методических материалов": 51,
        "промежуточная аттестация": 31,
    }
    assert (sem.credits, sem.active_hours, sem.total_hours) == (6, 62, 216)
    assert total.is_total
    assert not [w for w in result.warnings if w.startswith("трудоёмкость")]


def test_merged_cells_table_with_extra_columns(tmp_path: Path) -> None:
    """В части документов столбцов 21, а не 19: столбцы опознаются по смыслу заголовка."""
    text = (FIXTURES / "rpd_full_synthetic.md").read_text(encoding="utf-8")
    lines = text.splitlines()
    widened = [
        ln.replace(
            " |  |  |  |  |  |  |  |  |  |  |  |  | Самостоятельная",
            " |  |  |  |  |  |  |  |  |  |  |  |  |  | Самостоятельная",
            1,
        )
        if "Контактная работа" in ln
        else ln
        for ln in lines
    ]
    result = extract_rpd("\n".join(widened), file="x.md")
    assert result.workload_table  # лишний столбец не ломает разбор, в отличие от позиционного


@pytest.mark.parametrize(
    ("line", "expected"),
    [
        ("Практические занятия – 48 час. Решение задач по пособию.", [48]),
        ("Лекции, 8 ч Лабораторных работ, 50 ч. Самостоятельной работы, 24 ч.", [8, 50, 24]),
        ("Консультация (2 часа), экзамен в устной форме.", [2]),
        ("Объём составляет 32 часа: семинары — 16 часов, практические занятия — 16 часов.", [32]),
    ],
)
def test_active_forms_phrasings_from_real_documents(line: str, expected: list[int]) -> None:
    text = f"## 1.4. Перечень активных и интерактивных форм учебных занятий\n\n{line}\n"
    assert [f.hours for f in extract_rpd(text, "x.md").active_forms] == expected
